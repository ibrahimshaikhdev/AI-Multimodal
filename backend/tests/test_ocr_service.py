from io import BytesIO
from types import SimpleNamespace

import fitz
from PIL import Image
from docx import Document

from backend.services.ocr_service import extract_document_text, extract_image_text


def test_extract_image_text_returns_readable_text(monkeypatch):
    image = Image.new("RGB", (300, 120), color="white")
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    image_bytes = buffer.getvalue()

    def fake_image_to_string(image_obj, lang="eng"):
        assert lang == "eng"
        assert image_obj is not None
        return "Patient: Jordan Lee\nHemoglobin: 14.2 g/dL\nDoctor: Dr. Smith"

    monkeypatch.setattr(
        "backend.services.ocr_service.pytesseract",
        SimpleNamespace(image_to_string=fake_image_to_string),
    )

    extracted = extract_image_text(image_bytes)

    assert "Jordan Lee" in extracted
    assert "Hemoglobin" in extracted
    assert "Dr. Smith" in extracted


def test_extract_document_text_dispatches_for_images(monkeypatch):
    image = Image.new("RGB", (240, 80), color="white")
    buffer = BytesIO()
    image.save(buffer, format="JPEG")
    image_bytes = buffer.getvalue()

    def fake_image_to_string(image_obj, lang="eng"):
        return "Lab review\nBP: 118/76"

    monkeypatch.setattr(
        "backend.services.ocr_service.pytesseract",
        SimpleNamespace(image_to_string=fake_image_to_string),
    )

    extracted = extract_document_text(image_bytes, "scan.jpg")

    assert "Lab review" in extracted
    assert "118/76" in extracted


def test_extract_document_text_uses_ocr_for_scanned_pdf_pages(monkeypatch):
    document = fitz.open()
    document.new_page()
    pdf_bytes = document.tobytes()
    document.close()

    monkeypatch.setattr(
        "backend.services.ocr_service.extract_image_text",
        lambda image_bytes, language="eng": "MRI report: no acute finding",
    )

    extracted = extract_document_text(pdf_bytes, "scan.pdf")

    assert extracted == "MRI report: no acute finding"


def test_extract_document_text_reads_docx_paragraphs_and_tables():
    document = Document()
    document.add_paragraph("Patient: Jordan Lee")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Hemoglobin"
    table.cell(0, 1).text = "14.2 g/dL"
    buffer = BytesIO()
    document.save(buffer)

    extracted = extract_document_text(buffer.getvalue(), "report.docx")

    assert "Patient: Jordan Lee" in extracted
    assert "Hemoglobin | 14.2 g/dL" in extracted


def test_extract_document_text_rejects_unsupported_content():
    try:
        extract_document_text(b"not an image", "notes.txt")
        assert False, "Expected ValueError for unsupported document type"
    except ValueError:
        pass

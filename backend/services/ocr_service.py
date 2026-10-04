from __future__ import annotations

import os
import shutil
from io import BytesIO
from pathlib import Path

import fitz

from backend.services.ocr_preprocessing import preprocess_image_for_ocr

try:
    import pytesseract
except ImportError:  # pragma: no cover - optional dependency path
    pytesseract = None

try:
    from PIL import Image, UnidentifiedImageError
except ImportError:  # pragma: no cover - optional dependency path
    Image = None
    UnidentifiedImageError = OSError

SUPPORTED_IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"}


def _normalize_text(raw_text: str) -> str:
    return "\n".join(part.strip() for part in str(raw_text).splitlines() if part.strip()).strip()


def _configure_tesseract_path():
    if pytesseract is None:
        return

    if shutil.which("tesseract"):
        return

    candidate_paths = [
        r"C:\Program Files\Tesseract-OCR\tesseract.exe",
        r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
    ]

    for candidate in candidate_paths:
        if os.path.exists(candidate):
            if hasattr(pytesseract, "pytesseract"):
                pytesseract.pytesseract.tesseract_cmd = candidate
            os.environ["PATH"] = os.pathsep.join(
                [str(Path(candidate).parent), os.environ.get("PATH", "")]
            )
            return


def _ensure_ocr_ready():
    if Image is None or pytesseract is None:
        raise RuntimeError(
            "OCR dependencies are not installed. Install Pillow and pytesseract, and ensure Tesseract is available on PATH."
        )

    _configure_tesseract_path()


def _open_image(image_bytes: bytes):
    if not isinstance(image_bytes, (bytes, bytearray)) or not image_bytes:
        raise ValueError("Image content is required")

    try:
        image = Image.open(BytesIO(image_bytes))
    except (UnidentifiedImageError, OSError) as exc:
        raise ValueError("Invalid image content") from exc

    return image


def extract_image_text(image_bytes: bytes, language: str = "eng") -> str:
    _ensure_ocr_ready()

    processed_bytes = preprocess_image_for_ocr(image_bytes)
    image = _open_image(processed_bytes)
    try:
        try:
            raw_text = pytesseract.image_to_string(image, lang=language)
        except Exception as exc:  # pragma: no cover - OCR engine path is environment-specific
            raise RuntimeError(
                "Tesseract OCR failed to process the image. Check the Tesseract installation and PATH configuration."
            ) from exc
    finally:
        image.close()

    cleaned_text = _normalize_text(raw_text)
    if not cleaned_text:
        raise ValueError("No readable text found in the image")

    return cleaned_text


def extract_document_text(document_bytes: bytes, filename: str | None = None, language: str = "eng") -> str:
    if not isinstance(document_bytes, (bytes, bytearray)) or not document_bytes:
        raise ValueError("Document content is required")

    suffix = (Path(filename).suffix.lower() if filename else "").strip()

    if suffix == ".pdf":
        try:
            document = fitz.open(stream=bytes(document_bytes), filetype="pdf")
        except Exception as exc:
            raise ValueError("Invalid PDF content") from exc

        page_texts = []
        try:
            for page in document:
                page_text = page.get_text("text", sort=True).strip()
                if page_text:
                    page_texts.append(page_text)
                    continue

                page_image = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                try:
                    page_text = extract_image_text(page_image.tobytes("png"), language=language)
                except ValueError as exc:
                    if str(exc) != "No readable text found in the image":
                        raise
                    continue
                page_texts.append(page_text)
        finally:
            document.close()

        if not page_texts:
            raise ValueError("No readable text found in the document")
        return "\n\n".join(page_texts)

    if suffix in SUPPORTED_IMAGE_EXTENSIONS:
        return extract_image_text(document_bytes, language=language)

    if suffix == ".docx":
        try:
            from docx import Document

            document = Document(BytesIO(bytes(document_bytes)))
        except Exception as exc:
            raise ValueError("Invalid DOCX content") from exc

        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)

        if not paragraphs:
            raise ValueError("No readable text found in the document")
        return "\n".join(paragraphs)

    if not suffix:
        raise ValueError("File type is required to determine OCR extraction strategy")

    raise ValueError("Unsupported document type for OCR extraction")


__all__ = ["extract_image_text", "extract_document_text"]
